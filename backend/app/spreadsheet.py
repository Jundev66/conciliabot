from __future__ import annotations

import io
from typing import Any, Iterable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Spreadsheet support for ConciliaBot.
#
# The example/preview workbook mirrors the delivery-note schema that CoreBiz
# exposes on /v1/integrations/delivery-notes (DeliveryNoteListItem plus part of
# DeliveryNoteView):
#
#   number, customerName, currency, total, totalSecondary, status, issuedAt,
#   deliveredAt, receivedBy
#
# Keeping the header names and the "delivered => received" rule aligned with
# app/agent/corebiz.py lets a client validate the reconciliation offline with a
# file that has the same shape as the real integration.

DELIVERY_NOTE_COLUMNS: list[tuple[str, str]] = [
    ("number", "Nº nota de entrega"),
    ("customerName", "Cliente"),
    ("currency", "Moneda"),
    ("total", "Total"),
    ("totalSecondary", "Total (2ª moneda)"),
    ("status", "Estado"),
    ("issuedAt", "Emitida (ISO)"),
    ("deliveredAt", "Entregada (ISO)"),
    ("receivedBy", "Recibido por"),
]

INVOICE_COLUMNS: list[tuple[str, str]] = [
    ("folio", "Folio"),
    ("uuid", "UUID"),
    ("provider", "Proveedor"),
    ("taxId", "RFC/NIT"),
    ("country", "País"),
    ("total", "Total"),
    ("currency", "Moneda"),
    ("status", "Estado"),
    ("deliveryNoteNumber", "Nº nota de entrega"),
    ("deliveryNoteAmount", "Importe nota"),
    ("issueDate", "Emitida"),
    ("createdAt", "Registrada"),
]

# Accept both the Spanish headers we write and the raw CoreBiz field names.
_HEADER_TO_KEY: dict[str, str] = {header: key for key, header in DELIVERY_NOTE_COLUMNS}
_HEADER_TO_KEY.update({key: key for key, _ in DELIVERY_NOTE_COLUMNS})

_DELIVERED_STATUSES = {"entregada", "delivered", "recibida", "received"}

# Synthetic example rows. Customers are fictional and the amounts are invented;
# nothing here is a real document.
SAMPLE_ROWS: list[dict[str, Any]] = [
    {
        "number": "NE-5001",
        "customerName": "Distribuidora del Norte S.A. de C.V.",
        "currency": "MXN",
        "total": 14500.0,
        "totalSecondary": 0.0,
        "status": "entregada",
        "issuedAt": "2026-01-15",
        "deliveredAt": "2026-01-18",
        "receivedBy": "Juan Mata",
    },
    {
        "number": "NE-5002",
        "customerName": "Aceros y Perfiles Monterrey S.A.",
        "currency": "MXN",
        "total": 16000.0,
        "totalSecondary": 0.0,
        "status": "emitida",
        "issuedAt": "2026-01-16",
        "deliveredAt": "",
        "receivedBy": "",
    },
    {
        "number": "NE-5003",
        "customerName": "Metalurgica Paulista Ltda.",
        "currency": "BRL",
        "total": 9800.0,
        "totalSecondary": 0.0,
        "status": "anulada",
        "issuedAt": "2026-01-17",
        "deliveredAt": "",
        "receivedBy": "",
    },
]

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000


class SpreadsheetError(ValueError):
    """Raised when an uploaded workbook cannot be read as delivery notes."""


def _to_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_sheet(ws: Any, columns: list[tuple[str, str]], rows: Iterable[dict[str, Any]]) -> None:
    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(bold=True, color="FFFFFF")
    for col, (_key, header) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    for r, row in enumerate(rows, start=2):
        for col, (key, _header) in enumerate(columns, start=1):
            value = row.get(key)
            ws.cell(row=r, column=col, value="" if value is None else value)
    for col, (_key, header) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col)].width = max(len(str(header)) + 4, 14)


def build_delivery_notes_workbook(rows: Optional[list[dict[str, Any]]] = None) -> bytes:
    """Example workbook with the CoreBiz delivery-note columns."""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Notas de entrega"
    _write_sheet(ws, DELIVERY_NOTE_COLUMNS, SAMPLE_ROWS if rows is None else rows)
    return _to_bytes(workbook)


def build_invoices_workbook(invoices: Iterable[Any]) -> bytes:
    """Export the reconciled records (invoices) as a workbook."""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Registros"
    _write_sheet(ws, INVOICE_COLUMNS, [_invoice_row(inv) for inv in invoices])
    return _to_bytes(workbook)


def _invoice_row(inv: Any) -> dict[str, Any]:
    note = getattr(inv, "delivery_note", None)
    return {
        "folio": inv.folio,
        "uuid": inv.uuid,
        "provider": inv.provider.name,
        "taxId": inv.provider.tax_id,
        "country": inv.country.value,
        "total": inv.total,
        "currency": inv.currency,
        "status": inv.status.value,
        "deliveryNoteNumber": note.number if note else "",
        "deliveryNoteAmount": note.amount if note else "",
        "issueDate": inv.issue_date,
        "createdAt": inv.created_at,
    }


def _as_number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").replace("$", "").strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _normalize_row(item: dict[str, Any]) -> dict[str, Any]:
    total = _as_number(item.get("total"))
    status = str(item.get("status") or "").strip()
    delivered_at = item.get("deliveredAt")
    delivered = status.lower() in _DELIVERED_STATUSES or (
        delivered_at is not None and str(delivered_at).strip() != ""
    )
    return {
        "number": str(item.get("number") or "").strip(),
        "customerName": str(item.get("customerName") or "").strip(),
        "currency": str(item.get("currency") or "").strip(),
        "total": total if total is not None else 0.0,
        "totalSecondary": _as_number(item.get("totalSecondary")) or 0.0,
        "status": status,
        "issuedAt": str(item.get("issuedAt") or "").strip(),
        "deliveredAt": str(delivered_at or "").strip(),
        "receivedBy": str(item.get("receivedBy") or "").strip(),
        "receivedAmount": (total or 0.0) if delivered else 0.0,
    }


def parse_delivery_notes(data: bytes, *, max_rows: int = MAX_ROWS) -> dict[str, Any]:
    """Read an .xlsx of delivery notes and validate each row (read-only)."""
    if not data:
        raise SpreadsheetError("archivo vacío")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure is a bad upload
        raise SpreadsheetError("no parece un archivo .xlsx válido") from exc
    ws = workbook.active
    if ws is None:
        raise SpreadsheetError("el libro no tiene hojas")

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise SpreadsheetError("la hoja está vacía")

    keys = [_HEADER_TO_KEY.get(str(h).strip()) if h is not None else None for h in header]
    if not any(keys):
        raise SpreadsheetError("no se reconoce ninguna columna de nota de entrega")

    parsed: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    truncated = False
    for offset, raw in enumerate(rows_iter, start=2):
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        if len(parsed) >= max_rows:
            truncated = True
            break
        item: dict[str, Any] = {}
        for idx, key in enumerate(keys):
            if key is None or idx >= len(raw):
                continue
            item[key] = raw[idx]
        row = _normalize_row(item)
        if not row["number"]:
            errors.append({"row": offset, "field": "number", "message": "falta el número de nota"})
        if _as_number(item.get("total")) is None:
            errors.append({"row": offset, "field": "total", "message": "el total no es numérico"})
        if not row["currency"]:
            errors.append({"row": offset, "field": "currency", "message": "falta la moneda"})
        parsed.append(row)

    return {
        "headers": [str(h) for h in header if h is not None and str(h).strip() != ""],
        "count": len(parsed),
        "rows": parsed,
        "errors": errors,
        "truncated": truncated,
    }
